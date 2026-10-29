import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from sapagent import ai
from sapagent.classifier import classify_orders, classify_text
from sapagent.kpis import analyse, daily_brief
from sapagent.loader import ExportError, load_export
from sapagent.pipeline import run
from sapagent.report import build_daily_brief, build_report
from sapagent.sample_data import generate

SMALL_CSV = """AUFNR;AUART;EQUNR;EQKTX;KTEXT;ERDAT;AUSVN;AUSBS;MSAUS;STTXT;GKSTI
1;PM04;P1;Pompe;fuite garniture;01.01.2026;01.01.2026 08:00;01.01.2026 12:00;X;TECO;1.000,50
2;PM01;P1;Pompe;fuite garniture mécanique;15.01.2026;15.01.2026 08:00;15.01.2026 10:00;X;TECO;500
3;PM04;P1;Pompe;chgt garniture fuite;30.01.2026;30.01.2026 08:00;30.01.2026 14:00;X;TECO;250
4;PM02;P1;Pompe;Graissage hebdomadaire;10.01.2026;;;;TECO;100
5;PM01;C1;Convoyeur;voir chef de quart;20.01.2026;20.01.2026 08:00;20.01.2026 09:00;;REL;80
"""


def load_small():
    return load_export(SMALL_CSV.encode(), "export.csv")


def test_loader_maps_sap_technical_names_and_european_numbers():
    res = load_small()
    assert {"order_id", "order_type", "equipment_id", "malfunction_start", "actual_cost"} <= set(res.mapping)
    df = res.data
    assert df.loc[0, "actual_cost"] == pytest.approx(1000.50)
    assert df.loc[0, "downtime_h"] == pytest.approx(4.0)
    assert df.loc[4, "downtime_h"] == 0  # no breakdown flag -> not a stop
    assert df["is_failure"].tolist() == [True, True, True, False, True]
    assert df["is_open"].tolist() == [False, False, False, False, True]


def test_loader_accepts_french_alv_headers_from_sample():
    buf = io.BytesIO()
    generate(months=2).to_excel(buf, index=False)
    res = load_export(buf.getvalue(), "iw39.xlsx")
    assert not res.missing_recommended
    assert not res.unmapped


def test_loader_rejects_file_without_key_columns():
    with pytest.raises(ExportError):
        load_export(b"foo;bar\n1;2\n", "x.csv")


@pytest.mark.parametrize("text,expected", [
    ("Roulement HS côté accouplement", "bearing"),
    ("rlt chauffe palier", "bearing"),
    ("mechanical seal leak", "seal"),
    ("Bande déchirée", "belt"),
    ("flexible hydraulique éclaté", "hydraulic"),
    ("goulotte bouchée", "blockage"),
    ("voir chef de quart", "other"),
])
def test_keyword_classifier(text, expected):
    assert classify_text(text)[0] == expected


def test_kpis_on_hand_computed_example():
    res = load_small()
    a = analyse(classify_orders(res.data), res.missing_recommended)
    p1 = a.equipment.set_index("equipment_id").loc["P1"]
    assert p1["failures"] == 3
    assert p1["downtime_h"] == pytest.approx(12.0)
    assert p1["mttr_h"] == pytest.approx(4.0)
    assert p1["mtbf_h"] == pytest.approx((a.period_hours - 12) / 3)
    assert a.overview["planned_ratio"] == pytest.approx(0.2)
    # Three seal leaks on P1, two of them within 30 days of the previous one.
    assert len(a.repeats) == 1
    assert a.repeats.iloc[0]["category"] == "seal"


def test_sample_data_surfaces_the_planted_problems():
    r = run(generate().to_csv(sep=";", index=False).encode(), "demo.csv")
    a = r.analysis
    assert "Pompe à pulpe P-201" in a.equipment.head(3)["equipment_desc"].tolist()
    assert a.repeats.iloc[0]["failure_mode"] == "Étanchéité / Fuites"
    assert a.quality["duplicate_orders"] == 1
    assert a.quality["negative_durations"] == 1
    assert 0 < a.overview["overdue_orders"] <= a.overview["open_orders"]
    assert r.insights["source"] == "rules" and r.insights["recommendations"]


def test_report_is_a_valid_workbook_with_all_sheets():
    r = run(generate(months=3).to_csv(sep=";", index=False).encode(), "demo.csv")
    wb = load_workbook(io.BytesIO(build_report(r.analysis, r.insights)))
    assert wb.sheetnames[:3] == ["Synthèse", "Équipements", "Modes de défaillance"]
    assert wb["Équipements"]["D5"].value.startswith("=COUNTIFS(")


# ---------------------------------------------------------------- IW39 / IW29 specifics
IW29_ENGLISH = """Notification,Notif. type,Equipment,Description,Notif.date,Malfunct.start,Malfunct.end,Breakdown,System status,Main WorkCtr
900001,M2,P1,Seal leak on pump,2026-03-01,2026-03-01 08:00,2026-03-01 11:00,X,NOCO,HCNEMNT
900002,M1,P1,Request: paint guard,2026-03-02,,,,OSNO,HCSEMNT
900003,M2,C7,belt misalignment,2026-03-03,2026-03-03 10:00,,X,OSNO,HCWAMNT
"""


def test_notification_export_in_english_without_order_numbers():
    res = load_export(IW29_ENGLISH.encode(), "iw29.csv")
    df = res.data
    assert "order_id" not in res.mapping and df["order_id"].tolist() == ["900001", "900002", "900003"]
    assert df["order_kind"].tolist() == ["corrective", "request", "corrective"]
    assert df["is_open"].tolist() == [False, True, True]
    assert df.loc[0, "downtime_h"] == pytest.approx(3.0)
    assert df["work_center"].tolist() == ["HCNEMNT", "HCSEMNT", "HCWAMNT"]


def test_work_center_filter_accepts_sap_wildcards():
    assert len(run(IW29_ENGLISH.encode(), "iw29.csv", work_centers="HCNE*, HCWA*").analysis.orders) == 2
    assert len(run(IW29_ENGLISH.encode(), "iw29.csv", work_centers="hc*").analysis.orders) == 3
    with pytest.raises(ExportError):
        run(IW29_ENGLISH.encode(), "iw29.csv", work_centers="ZZZ*")


def test_overdue_orders_use_basic_start_date():
    csv = """Order,Order Type,Equipment,Description,Created on,Bsc start,System status,Main work ctr
1,PM02,E1,Visite,01.01.2026,05.01.2026,REL,MECA
2,PM02,E1,Visite,01.01.2026,31.01.2026,REL,MECA
3,PM01,E2,Fuite,10.01.2026,12.01.2026,TECO,ELEC
4,PM01,E2,Fuite,31.01.2026,31.01.2026,CRTD,ELEC
"""
    res = load_export(csv.encode(), "iw39.csv")
    a = analyse(classify_orders(res.data))
    late = a.backlog.set_index("order_id")
    assert bool(late.at["1", "is_overdue"]) and late.at["1", "days_late"] == 26
    assert not late.at["2", "is_overdue"] and not late.at["4", "is_overdue"]
    wc = a.workcenters.set_index("work_center")
    assert wc.at["MECA", "overdue"] == 1 and wc.at["ELEC", "open_orders"] == 1
    assert a.overview["overdue_orders"] == 1


def test_daily_brief_lists_the_day_and_machines_still_down():
    r = run(generate().to_csv(sep=";", index=False).encode(), "demo.csv")
    brief = daily_brief(r.analysis)
    assert brief["day"] == r.analysis.period_end - pd.Timedelta(days=1)
    assert brief["n_failures"] == 4 and brief["n_stops"] == 3 and brief["n_recurring"] == 2
    assert brief["still_down"]["equipment_desc"].tolist() == ["Pompe à pulpe P-201"]
    wb = load_workbook(io.BytesIO(build_daily_brief(brief)))
    assert wb.sheetnames == ["Réunion J-1"]
    assert "31/08/2026" in wb["Réunion J-1"]["A1"].value


# ---------------------------------------------------------------- AI layer (no network)
def test_ai_classification_only_overrides_uncertain_rows(monkeypatch):
    res = load_small()
    df = classify_orders(res.data)

    def fake_ask(cfg, system, prompt, schema, effort, name):
        assert "voir chef de quart" in prompt
        assert "fuite garniture" not in prompt  # confident rule hits are not sent
        return {"items": [{"id": 0, "category": "electrical", "confidence": 0.7}]}

    monkeypatch.setattr(ai, "_ask_json", fake_ask)
    out = ai.classify_with_ai(df, ai.AIConfig("anthropic", "sk-test"))
    row = out[out["description"] == "voir chef de quart"].iloc[0]
    assert row["category"] == "electrical" and row["category_source"] == "IA (Claude)"


def test_pipeline_falls_back_to_rules_when_ai_fails(monkeypatch):
    def boom(*args, **kwargs):
        raise RuntimeError("réseau indisponible")

    monkeypatch.setattr(ai, "classify_with_ai", boom)
    r = run(SMALL_CSV.encode(), "export.csv", ai_config=ai.AIConfig("openrouter", "sk-or-test", "x/y"))
    assert r.insights["source"] == "rules"
    assert "réseau indisponible" in r.ai_error


class FakeResponse:
    def __init__(self, status, payload):
        self.status_code, self._payload, self.text = status, payload, str(payload)

    def json(self):
        return self._payload


def test_openrouter_request_uses_strict_json_schema(monkeypatch):
    calls = []

    def fake_post(url, headers, json, timeout):
        calls.append((url, headers, json))
        return FakeResponse(200, {"choices": [{"finish_reason": "stop", "message": {"content": '{"ok": true}'}}]})

    monkeypatch.setattr(ai.requests, "post", fake_post)
    cfg = ai.AIConfig("openrouter", "sk-or-test", "anthropic/claude-sonnet-4.5")
    assert ai._ask_json(cfg, "sys", "prompt", {"type": "object"}, "low", "test") == {"ok": True}
    url, headers, body = calls[0]
    assert url == "https://openrouter.ai/api/v1/chat/completions"
    assert headers["Authorization"] == "Bearer sk-or-test"
    assert body["model"] == "anthropic/claude-sonnet-4.5"
    assert body["messages"][0] == {"role": "system", "content": "sys"}
    assert body["response_format"]["type"] == "json_schema"
    assert body["response_format"]["json_schema"]["strict"] is True
    assert body["provider"] == {"require_parameters": True}


def test_openrouter_falls_back_to_json_mode_for_models_without_structured_outputs(monkeypatch):
    calls = []

    def fake_post(url, headers, json, timeout):
        calls.append(json)
        if len(calls) == 1:
            return FakeResponse(404, {"error": {"code": 404, "message": "No endpoints found that support the requested parameters"}})
        return FakeResponse(200, {"choices": [{"message": {"content": 'Voici :\n```json\n{"items": []}\n```'}}]})

    monkeypatch.setattr(ai.requests, "post", fake_post)
    cfg = ai.AIConfig("openrouter", "sk-or-test", "meta-llama/some-model")
    assert ai._ask_json(cfg, "sys", "prompt", {"type": "object"}, "low", "test") == {"items": []}
    assert calls[1]["response_format"] == {"type": "json_object"} and "provider" not in calls[1]
    assert "schéma JSON" in calls[1]["messages"][-1]["content"]


def test_openrouter_errors_are_reported(monkeypatch):
    monkeypatch.setattr(ai.requests, "post",
                        lambda *a, **k: FakeResponse(401, {"error": {"code": 401, "message": "No auth credentials found"}}))
    with pytest.raises(ai.AIError, match="No auth credentials"):
        ai._ask_json(ai.AIConfig("openrouter", "bad", "x/y"), "s", "p", {}, "low", "t")


def test_ai_insights_are_sanitised_and_labelled(monkeypatch):
    res = load_small()
    a = analyse(classify_orders(res.data))
    monkeypatch.setattr(ai, "_ask_json", lambda *args, **kw: {
        "executive_summary": "Résumé.", "key_findings": ["A"],
        "recommendations": [{"priority": "Urgent", "equipment": "P1", "action": "RCA", "rationale": "3 pannes"},
                            {"priority": "Haute"}]})
    out = ai.ai_insights(a, ai.AIConfig("openrouter", "k", "openai/gpt-x"))
    assert out["source"] == "ai" and out["model"] == "openai/gpt-x"
    assert out["recommendations"] == [{"priority": "Moyenne", "equipment": "P1", "action": "RCA", "rationale": "3 pannes"}]


def test_default_model_prefers_newest_claude_opus():
    ids = ["openai/gpt-x", "anthropic/claude-sonnet-9", "anthropic/claude-opus-9", "anthropic/claude-opus-8"]
    assert ai.pick_default_model(ids) == "anthropic/claude-opus-9"
    assert ai.pick_default_model(["mistralai/m"]) == "mistralai/m"


def test_provider_is_picked_from_available_key(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("SAPAGENT_PROVIDER", raising=False)
    monkeypatch.setenv("OPENROUTER_API_KEY", "sk-or")
    cfg = ai.AIConfig.from_env()
    assert cfg.provider == "openrouter" and cfg.key == "sk-or"
